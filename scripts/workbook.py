"""Read and write the job tracker workbook.

The Applications sheet is the source of truth. Active Pipeline and Company
History are derived views that ship with formulas hardcoded to fixed rows, so
they are regenerated from scratch on every write rather than appended to.
"""
from __future__ import annotations

import copy
import datetime as dt
import re
import shutil
from collections import Counter, OrderedDict
from pathlib import Path

from openpyxl import load_workbook

from common import BACKUPS, CLOSED_STATUSES, STALE_AFTER_MONTHS, months_ago, thread_key

APPLICATIONS = "Applications"
ACTIVE_PIPELINE = "Active Pipeline"
COMPANY_HISTORY = "Company History"
LISTS = "Lists"

# Column letters in Applications, by field name.
COLS = OrderedDict([
    ("app_id", "A"), ("company", "B"), ("role", "C"), ("track", "D"),
    ("job_req", "E"), ("date_applied", "F"), ("source", "G"), ("status", "H"),
    ("status_detail", "I"), ("last_activity_date", "J"), ("attempt", "K"),
    ("prior_attempt", "L"), ("notes", "M"), ("days_since", "N"), ("thread_id", "O"),
])

DAYS_SINCE_FORMULA = '=IF(J{row}="","",TODAY()-J{row})'


def open_workbook(path: Path):
    return load_workbook(path)


def backup(path: Path) -> Path:
    BACKUPS.mkdir(exist_ok=True)
    stamp = dt.datetime.now().strftime("%Y-%m-%d_%H%M%S")
    dest = BACKUPS / f"{path.stem} {stamp}.xlsx"
    shutil.copy2(path, dest)
    return dest


def read_lists(wb) -> dict[str, list[str]]:
    """Controlled vocabularies: column A statuses, B tracks, C sources."""
    ws = wb[LISTS]
    out = {"status": [], "track": [], "source": []}
    for col, key in (("A", "status"), ("B", "track"), ("C", "source")):
        for row in range(1, ws.max_row + 1):
            val = ws[f"{col}{row}"].value
            if val and str(val).strip():
                out[key].append(str(val).strip())
    return out


def _as_date(value):
    if isinstance(value, dt.datetime):
        return value.date()
    if isinstance(value, dt.date):
        return value
    if isinstance(value, (int, float)):
        # Excel serial, 1900 date system with the historical leap-year bug.
        return dt.date(1899, 12, 30) + dt.timedelta(days=int(value))
    if isinstance(value, str) and value.strip():
        try:
            return dt.date.fromisoformat(value.strip()[:10])
        except ValueError:
            return None
    return None


def read_applications(wb) -> list[dict]:
    """Every populated Applications row, with its 1-based sheet row number."""
    ws = wb[APPLICATIONS]
    rows = []
    for r in range(2, ws.max_row + 1):
        app_id = ws[f"A{r}"].value
        if not app_id or not str(app_id).strip():
            continue
        rec = {"row": r}
        for field, col in COLS.items():
            rec[field] = ws[f"{col}{r}"].value
        rec["date_applied"] = _as_date(rec["date_applied"])
        rec["last_activity_date"] = _as_date(rec["last_activity_date"])
        rec["app_id"] = str(rec["app_id"]).strip()
        rec["company"] = (str(rec["company"]).strip() if rec["company"] else "")
        rec["status"] = (str(rec["status"]).strip() if rec["status"] else "")
        rows.append(rec)
    return rows


def known_thread_keys(apps: list[dict]) -> set[str]:
    keys = set()
    for a in apps:
        for part in str(a.get("thread_id") or "").split("/"):
            k = thread_key(part)
            if k:
                keys.add(k)
    return keys


def next_app_id(apps: list[dict]) -> int:
    highest = 0
    for a in apps:
        m = re.fullmatch(r"APP-(\d+)", a["app_id"])
        if m:
            highest = max(highest, int(m.group(1)))
    return highest + 1


def _copy_style(ws, src_row: int, dst_row: int) -> None:
    for col in COLS.values():
        src = ws[f"{col}{src_row}"]
        dst = ws[f"{col}{dst_row}"]
        if src.has_style:
            dst._style = copy.copy(src._style)


def apply_decisions(wb, decisions: list[dict], logger) -> dict:
    """Append new applications and patch existing ones. Returns a summary."""
    ws = wb[APPLICATIONS]
    apps = read_applications(wb)
    by_id = {a["app_id"]: a for a in apps}
    seen = known_thread_keys(apps)
    template_row = apps[-1]["row"] if apps else 1
    write_row = (apps[-1]["row"] + 1) if apps else 2
    counter = next_app_id(apps)
    added, updated, skipped = 0, 0, 0

    for d in decisions:
        action = d.get("action")
        tid = str(d.get("thread_id") or "").strip()
        key = thread_key(tid)

        if action == "ignore":
            skipped += 1
            continue

        if key and key in seen and action == "new_application":
            logger.info(f"  skip (already tracked): thread {tid} {d.get('company')}")
            skipped += 1
            continue

        if action == "update_existing":
            target = by_id.get(str(d.get("app_id") or "").strip())
            if not target:
                logger.warning(f"  update targets unknown app_id {d.get('app_id')!r}; treating as new")
                action = "new_application"
            else:
                r = target["row"]
                if d.get("status"):
                    ws[f"H{r}"] = d["status"]
                if d.get("status_detail"):
                    ws[f"I{r}"] = d["status_detail"]
                last = _as_date(d.get("last_activity_date"))
                if last:
                    ws[f"J{r}"] = dt.datetime(last.year, last.month, last.day)
                if d.get("notes"):
                    existing = str(ws[f"M{r}"].value or "").strip()
                    note = d["notes"].strip()
                    if note and note not in existing:
                        ws[f"M{r}"] = f"{existing}; {note}".strip("; ")
                if key and key not in seen:
                    prior = str(ws[f"O{r}"].value or "").strip()
                    ws[f"O{r}"] = f"{prior}/{tid}".strip("/")
                    seen.add(key)
                ws[f"N{r}"] = DAYS_SINCE_FORMULA.format(row=r)
                updated += 1
                logger.info(f"  update {target['app_id']}: {target['company']} -> {d.get('status')}")
                continue

        if action == "new_application":
            company = (d.get("company") or "").strip()
            if not company:
                logger.warning(f"  drop new application with no company (thread {tid})")
                skipped += 1
                continue

            prior_rows = [a for a in apps if a["company"].lower() == company.lower()]
            app_id = f"APP-{counter:03d}"
            counter += 1
            r = write_row
            write_row += 1

            if template_row > 1:
                _copy_style(ws, template_row, r)

            applied = _as_date(d.get("date_applied"))
            last = _as_date(d.get("last_activity_date")) or applied

            ws[f"A{r}"] = app_id
            ws[f"B{r}"] = company
            ws[f"C{r}"] = d.get("role") or ""
            ws[f"D{r}"] = d.get("track") or ""
            ws[f"E{r}"] = d.get("job_req") or ""
            if applied:
                ws[f"F{r}"] = dt.datetime(applied.year, applied.month, applied.day)
            ws[f"G{r}"] = d.get("source") or ""
            ws[f"H{r}"] = d.get("status") or ""
            ws[f"I{r}"] = d.get("status_detail") or ""
            if last:
                ws[f"J{r}"] = dt.datetime(last.year, last.month, last.day)
            ws[f"K{r}"] = len(prior_rows) + 1
            ws[f"L{r}"] = prior_rows[-1]["app_id"] if prior_rows else ""
            ws[f"M{r}"] = d.get("notes") or ""
            ws[f"N{r}"] = DAYS_SINCE_FORMULA.format(row=r)
            ws[f"O{r}"] = tid

            record = {"row": r, "app_id": app_id, "company": company,
                      "status": d.get("status") or "", "track": d.get("track") or "",
                      "date_applied": applied, "last_activity_date": last,
                      "source": d.get("source") or "", "thread_id": tid}
            apps.append(record)
            by_id[app_id] = record
            if key:
                seen.add(key)
            added += 1
            logger.info(f"  add {app_id}: {company} - {d.get('role')} [{d.get('status')}]")

    return {"added": added, "updated": updated, "skipped": skipped}


def expire_stale(wb, logger, today: dt.date | None = None) -> int:
    """Close Awaiting Response rows with no activity for STALE_AFTER_MONTHS.

    Measured from Last Activity Date, falling back to Date Applied, so any
    acknowledgement or follow-up restarts the clock. Returns the number closed.
    """
    ws = wb[APPLICATIONS]
    cutoff = months_ago(STALE_AFTER_MONTHS, today)
    detail = f"No response after {STALE_AFTER_MONTHS} months (auto-closed)"
    expired = 0
    for a in read_applications(wb):
        if a["status"] != "Awaiting Response":
            continue
        last = a["last_activity_date"] or a["date_applied"]
        if not last or last > cutoff:
            continue
        r = a["row"]
        ws[f"H{r}"] = "Rejected"
        ws[f"I{r}"] = detail
        expired += 1
        logger.info(f"  expire {a['app_id']}: {a['company']} (no activity since {last})")
    return expired


def regenerate_active_pipeline(wb) -> int:
    """Rebuild the view so it covers every active row, not a fixed 23."""
    ws = wb[ACTIVE_PIPELINE]
    apps = read_applications(wb)
    active = [a for a in apps if a["status"] not in CLOSED_STATUSES]
    active.sort(key=lambda a: (a["date_applied"] or dt.date.min), reverse=True)

    notes, _ = _take_notes(ws)
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for i, a in enumerate(active):
        out = i + 2
        src = a["row"]
        ws[f"A{out}"] = f"={APPLICATIONS}!B{src}"
        ws[f"B{out}"] = f"={APPLICATIONS}!C{src}"
        ws[f"C{out}"] = f"={APPLICATIONS}!D{src}"
        ws[f"D{out}"] = f"={APPLICATIONS}!H{src}"
        ws[f"E{out}"] = f'=IF({APPLICATIONS}!F{src}="","",TEXT({APPLICATIONS}!F{src},"yyyy-mm-dd"))'
        ws[f"F{out}"] = f'=IF({APPLICATIONS}!N{src}="","",{APPLICATIONS}!N{src})'
        ws[f"G{out}"] = f'=IF({APPLICATIONS}!M{src}="","",{APPLICATIONS}!M{src})'
    _put_notes(ws, notes, len(active) + 3, width=7)
    return len(active)


def _take_notes(ws) -> tuple[list[str], set[int]]:
    """Lift full-width merged note rows out of a derived sheet and unmerge it.

    Both derived sheets end with a merged explanatory row. delete_rows() does not
    clear merged ranges, and writing to a non-anchor MergedCell is silently
    discarded when the workbook is saved - which quietly ate a data row the first
    time this ran. So the ranges come out before any rewriting and the note text
    is re-appended underneath the regenerated rows.
    """
    notes, rows = [], set()
    for rng in list(ws.merged_cells.ranges):
        anchor = ws.cell(row=rng.min_row, column=rng.min_col)
        if anchor.value and str(anchor.value).strip():
            notes.append(str(anchor.value))
            rows.add(rng.min_row)
        ws.unmerge_cells(str(rng))
    return notes, rows


def _put_notes(ws, notes: list[str], start_row: int, width: int) -> None:
    for i, note in enumerate(notes):
        r = start_row + i
        ws.cell(row=r, column=1, value=note)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=width)


def _outcome_summary(rows: list[dict]) -> str:
    counts = Counter(r["status"] for r in rows if r["status"])
    if not counts:
        return ""
    return ", ".join(f"{n} {status.lower()}" for status, n in counts.most_common())


def regenerate_company_history(wb) -> int:
    """Rebuild per-company rows, preserving the hand-written prose columns."""
    ws = wb[COMPANY_HISTORY]
    apps = read_applications(wb)

    notes, note_rows = _take_notes(ws)

    # Preserve prose written by a human (or the original agent), keyed by company.
    preserved: dict[str, dict[str, str]] = {}
    for r in range(2, ws.max_row + 1):
        if r in note_rows:
            continue
        name = ws[f"A{r}"].value
        if not name or not str(name).strip():
            continue
        preserved[str(name).strip().lower()] = {
            "F": ws[f"F{r}"].value, "G": ws[f"G{r}"].value, "H": ws[f"H{r}"].value,
        }

    grouped: dict[str, list[dict]] = OrderedDict()
    for a in apps:
        if a["company"]:
            grouped.setdefault(a["company"], []).append(a)

    companies = sorted(grouped.items(), key=lambda kv: (-len(kv[1]), kv[0].lower()))

    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for i, (company, rows) in enumerate(companies):
        out = i + 2
        prior = preserved.get(company.lower(), {})
        tracks = sorted({r["track"] for r in rows if r.get("track")})
        ws[f"A{out}"] = company
        ws[f"B{out}"] = f'=COUNTIF({APPLICATIONS}!$B$2:$B$500,A{out})'
        ws[f"C{out}"] = ", ".join(str(t) for t in tracks)
        ws[f"D{out}"] = (f'=IFERROR(TEXT(_xlfn.MINIFS({APPLICATIONS}!$F$2:$F$500,'
                         f'{APPLICATIONS}!$B$2:$B$500,A{out}),"yyyy-mm-dd"),"")')
        ws[f"E{out}"] = (f'=IFERROR(TEXT(_xlfn.MAXIFS({APPLICATIONS}!$J$2:$J$500,'
                         f'{APPLICATIONS}!$B$2:$B$500,A{out}),"yyyy-mm-dd"),"")')
        # Regenerate the outcome tally for new companies; keep existing prose.
        ws[f"F{out}"] = prior.get("F") or _outcome_summary(rows)
        ws[f"G{out}"] = prior.get("G") or ""
        ws[f"H{out}"] = prior.get("H") or ""
    _put_notes(ws, notes, len(companies) + 3, width=8)
    return len(companies)
